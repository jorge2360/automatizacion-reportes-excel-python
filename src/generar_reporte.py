import pandas as pd
from pathlib import Path
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, Reference

RUTA_ENTRADA = Path("entrada")
RUTA_SALIDA = Path("salida")


def aplicar_formato_excel(ruta_archivo):
    encabezado_fill = PatternFill("solid", fgColor="1F4E78")
    encabezado_font = Font(color="FFFFFF", bold=True)
    borde = Border(
        left=Side(style="thin", color="D9E2F3"),
        right=Side(style="thin", color="D9E2F3"),
        top=Side(style="thin", color="D9E2F3"),
        bottom=Side(style="thin", color="D9E2F3"),
    )

    from openpyxl import load_workbook

    libro = load_workbook(ruta_archivo)

    for hoja in libro.worksheets:
        hoja.freeze_panes = "A2"

        for celda in hoja[1]:
            celda.fill = encabezado_fill
            celda.font = encabezado_font
            celda.alignment = Alignment(horizontal="center")
            celda.border = borde

        for fila in hoja.iter_rows(min_row=2):
            for celda in fila:
                celda.border = borde
                celda.alignment = Alignment(vertical="center")

        for columna in hoja.columns:
            max_length = 0
            letra_columna = get_column_letter(columna[0].column)

            for celda in columna:
                if celda.value:
                    max_length = max(max_length, len(str(celda.value)))

                if "precio" in str(hoja.cell(row=1, column=celda.column).value).lower() \
                    or "total" in str(hoja.cell(row=1, column=celda.column).value).lower() \
                    or "valor" in str(hoja.cell(row=1, column=celda.column).value).lower():
                    celda.number_format = 'Q #,##0.00'

            hoja.column_dimensions[letra_columna].width = max_length + 4

    libro.save(ruta_archivo)

def agregar_graficos_excel(ruta_archivo):
    from openpyxl import load_workbook

    libro = load_workbook(ruta_archivo)

    hoja_categoria = libro["Ventas por categoria"]

    grafico_categoria = BarChart()
    grafico_categoria.title = "Ventas por categoría"
    grafico_categoria.y_axis.title = "Total vendido"
    grafico_categoria.x_axis.title = "Categoría"

    datos = Reference(hoja_categoria, min_col=2, min_row=1, max_row=hoja_categoria.max_row)
    categorias = Reference(hoja_categoria, min_col=1, min_row=2, max_row=hoja_categoria.max_row)

    grafico_categoria.add_data(datos, titles_from_data=True)
    grafico_categoria.set_categories(categorias)

    hoja_categoria.add_chart(grafico_categoria, "D2")

    hoja_producto = libro["Ventas por producto"]

    grafico_producto = BarChart()
    grafico_producto.title = "Ventas por producto"
    grafico_producto.y_axis.title = "Total vendido"
    grafico_producto.x_axis.title = "Producto"

    datos_producto = Reference(hoja_producto, min_col=2, min_row=1, max_row=hoja_producto.max_row)
    productos = Reference(hoja_producto, min_col=1, min_row=2, max_row=hoja_producto.max_row)

    grafico_producto.add_data(datos_producto, titles_from_data=True)
    grafico_producto.set_categories(productos)

    hoja_producto.add_chart(grafico_producto, "D2")

    libro.save(ruta_archivo)

def crear_resumen_ejecutivo(ventas, inventario):
    total_ventas = ventas["total"].sum()
    total_unidades_vendidas = ventas["cantidad"].sum()
    total_productos = inventario["producto"].nunique()
    valor_total_inventario = inventario["valor_inventario"].sum()

    resumen = pd.DataFrame({
        "Indicador": [
            "Total de ventas",
            "Unidades vendidas",
            "Productos registrados",
            "Valor total del inventario",
        ],
        "Valor": [
            total_ventas,
            total_unidades_vendidas,
            total_productos,
            valor_total_inventario,
        ],
    })

    return resumen

def generar_reporte():
    RUTA_SALIDA.mkdir(exist_ok=True)

    ventas = pd.read_csv(RUTA_ENTRADA / "ventas.csv")
    inventario = pd.read_csv(RUTA_ENTRADA / "inventario.csv")

    ventas["total"] = ventas["cantidad"] * ventas["precio_unitario"]
    inventario["valor_inventario"] = inventario["stock"] * inventario["precio_compra"]

    ventas_por_categoria = ventas.groupby("categoria")["total"].sum().reset_index()
    ventas_por_producto = ventas.groupby("producto")["total"].sum().reset_index()
    inventario_valorizado = inventario[["producto", "categoria", "stock", "valor_inventario"]]
    resumen_ejecutivo = crear_resumen_ejecutivo(ventas, inventario)

    ruta_reporte = RUTA_SALIDA / "reporte_empresarial.xlsx"

    with pd.ExcelWriter(ruta_reporte, engine="openpyxl") as writer:
        resumen_ejecutivo.to_excel(writer, sheet_name="Resumen ejecutivo", index=False)
        ventas.to_excel(writer, sheet_name="Ventas", index=False)
        inventario.to_excel(writer, sheet_name="Inventario", index=False)
        ventas_por_categoria.to_excel(writer, sheet_name="Ventas por categoria", index=False)
        ventas_por_producto.to_excel(writer, sheet_name="Ventas por producto", index=False)
        inventario_valorizado.to_excel(writer, sheet_name="Inventario valorizado", index=False)

    aplicar_formato_excel(ruta_reporte)
    agregar_graficos_excel(ruta_reporte)
    print(f"Reporte generado correctamente: {ruta_reporte}")


if __name__ == "__main__":
    generar_reporte()